from __future__ import annotations

import io

from app.ingestion.parser.base import EmptyMarkdownError, ParsedDocument, UnsupportedFileError

PDF_CONTENT_TYPE = "application/pdf"
DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_converter = None


def _get_converter():
    global _converter
    if _converter is None:
        from docling.document_converter import DocumentConverter

        _converter = DocumentConverter()
    return _converter


class DoclingParser:
    version = "docling-1"

    def parse(self, data: bytes, *, filename: str) -> ParsedDocument:
        try:
            from docling.datamodel.base_models import DocumentStream
        except ImportError as exc:  # pragma: no cover - dependency guard
            raise UnsupportedFileError(
                "Docling is not installed", details={"parser": "docling"}
            ) from exc

        source = DocumentStream(name=filename, stream=io.BytesIO(data))
        try:
            result = _get_converter().convert(source)
        except Exception as exc:  # noqa: BLE001 - corrupt / unreadable payloads
            raise EmptyMarkdownError("Could not parse the document with Docling") from exc

        markdown = (result.document.export_to_markdown() or "").strip()
        if not markdown:
            raise EmptyMarkdownError("The document has no extractable text (it may be scanned)")
        return ParsedDocument(markdown=markdown, page_offsets=None, parser_version=self.version)


class XlsxParser:
    version = "xlsx-1"

    def parse(self, data: bytes, *, filename: str) -> ParsedDocument:
        try:
            import openpyxl
        except ImportError as exc:  # pragma: no cover - dependency guard
            raise UnsupportedFileError(
                "openpyxl is not installed", details={"content_type": XLSX_CONTENT_TYPE}
            ) from exc

        try:
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 - corrupt / non-xlsx payloads
            raise EmptyMarkdownError("Could not read the spreadsheet") from exc

        lines: list[str] = []
        for ws in wb.worksheets:
            lines.append(f"## {ws.title}")
            for row in ws.iter_rows(values_only=True):
                cells = ["" if c is None else str(c) for c in row]
                if any(cell.strip() for cell in cells):
                    lines.append(" | ".join(cells))
        wb.close()

        markdown = "\n".join(lines).strip()
        if not markdown:
            raise EmptyMarkdownError("The spreadsheet has no data")
        return ParsedDocument(markdown=markdown, page_offsets=None, parser_version=self.version)
